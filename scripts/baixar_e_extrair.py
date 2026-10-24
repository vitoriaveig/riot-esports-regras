#!/usr/bin/env python3
"""Baixa os documentos listados em catalogo/catalogo.csv e extrai o texto para textos/<id>.md.

- Um arquivo por documento único (cartões que apontam para o mesmo arquivo compartilham o texto).
- PDFs: pdftotext (poppler) com marcação de página  <!-- página N -->  para permitir citação.
- DOCX: python-docx.  Google Docs: exportação em texto.  Google Sheets publicadas: CSV por aba em dados/.
- Gera manifest.csv (status, sha256, páginas, palavras, idioma, data de download) e secoes.csv
  (índice de seções numeradas: documento, número, título, página), útil para buscas rápidas.
- Por padrão os originais (PDF/DOCX) NÃO são versionados no Git (pasta originais/ no .gitignore);
  apenas o texto extraído, para fins de pesquisa, com link para a fonte oficial.

Uso:
  python scripts/baixar_e_extrair.py               # baixa só o que ainda não foi baixado
  python scripts/baixar_e_extrair.py --forcar      # baixa tudo de novo
  python scripts/baixar_e_extrair.py --id <id>     # um documento
  python scripts/baixar_e_extrair.py --local pasta # usa PDFs já baixados manualmente (nome = id ou hash)
"""
import argparse, csv, hashlib, io, json, re, shutil, subprocess, sys, tempfile, time
from datetime import datetime, timezone
from pathlib import Path

import requests

UA = "Mozilla/5.0 (pesquisa academica; arquivo de regulamentos de esports; contato via GitHub)"
PAUSA = 1.5  # segundos entre downloads
CAMPOS_MANIFEST = ["texto_id", "chave_arquivo", "titulo", "url", "url_download", "status", "detalhe", "http",
                   "sha256", "bytes", "paginas", "palavras", "idioma", "baixado_em", "arquivo_texto"]


# ---------------------------------------------------------------- utilidades
def agora():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def url_download(url):
    """Converte links de visualização em links de download direto quando possível."""
    m = re.search(r"docs\.google\.com/document/(?:u/\d+/)?d/e/([\w-]+)/pub", url)
    if m: return f"https://docs.google.com/document/d/e/{m.group(1)}/pub?output=txt", "gdoc_pub"
    m = re.search(r"docs\.google\.com/document/(?:u/\d+/)?d/([\w-]+)", url)
    if m: return f"https://docs.google.com/document/d/{m.group(1)}/export?format=txt", "gdoc"
    m = re.search(r"docs\.google\.com/spreadsheets/(?:u/\d+/)?d/(e/)?([\w-]+)", url)
    if m:
        base = f"https://docs.google.com/spreadsheets/d/{'e/' if m.group(1) else ''}{m.group(2)}"
        return base + "/pubhtml", "gsheet"
    m = re.search(r"drive\.google\.com/file/d/([\w-]+)", url)
    if m: return f"https://drive.usercontent.google.com/download?id={m.group(1)}&export=download&confirm=t", "drive"
    m = re.search(r"%22ids%22%3A%5B%22([\w-]+)%22", url)
    if m: return f"https://drive.usercontent.google.com/download?id={m.group(1)}&export=download&confirm=t", "drive"
    if "slack.com" in url: return None, "indisponivel"
    if url.lower().endswith(".docx"): return url, "docx"
    return url, "pdf"


def detectar_idioma(txt):
    amostra = txt[:20000]
    if not amostra.strip(): return ""
    hangul = len(re.findall(r"[가-힣]", amostra)); kana = len(re.findall(r"[぀-ヿ]", amostra))
    cjk = len(re.findall(r"[一-鿿]", amostra)); letras = len(re.findall(r"[A-Za-zÀ-ÿ]", amostra)) or 1
    if hangul > letras * 0.2: return "ko"
    if kana > letras * 0.1: return "ja"
    if cjk > letras * 0.2: return "zh"
    palavras = re.findall(r"[a-záéíóúâêôãõçñü]+", amostra.lower())
    sw = {"en": {"the", "and", "of", "to", "shall", "be", "team", "player", "will", "must"},
          "pt": {"de", "e", "o", "a", "os", "das", "do", "da", "deve", "equipe", "jogador", "não"},
          "es": {"el", "los", "las", "del", "equipo", "jugador", "deberá", "y", "por"},
          "fr": {"le", "les", "des", "du", "équipe", "joueur", "doit", "et"},
          "de": {"der", "die", "und", "das", "spieler", "muss", "mannschaft"}}
    cont = {k: sum(1 for w in palavras if w in v) for k, v in sw.items()}
    return max(cont, key=cont.get) if max(cont.values()) > 20 else ""


RX_SECAO = re.compile(r"^\s{0,6}((?:\d{1,2}\.){1,4}\d{0,2}|[IVX]{1,5}\.|(?:Section|Seção|Artigo|Article|Rule|Regra)\s+\d+(?:\.\d+)*)\s+([A-ZÀ-Ý][^\n]{2,90})$")


def indexar_secoes(paginas):
    out = []
    for n, pg in enumerate(paginas, 1):
        for linha in pg.splitlines():
            m = RX_SECAO.match(linha.rstrip())
            if m and not re.search(r"(\.{4,}|\s{3,})\s*\d{1,3}\s*$", linha):  # ignora linhas de sumário
                out.append((m.group(1).rstrip("."), re.sub(r"\s{2,}", " ", m.group(2)).strip(), n))
    return out


RX_INVISIVEIS = re.compile("[\u00ad\u200b-\u200f\u202a-\u202e\u2060-\u2069\ufeff]")


def normalizar(txt):
    """Remove caracteres invisíveis (marcas de direção, espaços de largura zero) que quebram buscas."""
    return RX_INVISIVEIS.sub("", txt).replace("\u00a0", " ")


def links_pdf(dados):
    try:
        from pypdf import PdfReader
        out = []
        for pg in PdfReader(io.BytesIO(dados)).pages:
            for an in pg.get("/Annots") or []:
                a = an.get_object().get("/A") or {}
                u = a.get("/URI") if hasattr(a, "get") else None
                if u and str(u) not in out: out.append(str(u))
        return out
    except Exception:
        return []


def ocr(dados):
    """OCR (ocrmypdf + tesseract) para PDFs que são imagem. Retorna os bytes do PDF com camada de texto."""
    if not shutil.which("ocrmypdf"): return None
    with tempfile.TemporaryDirectory() as d:
        ent, sai = Path(d) / "e.pdf", Path(d) / "s.pdf"
        ent.write_bytes(dados)
        p = subprocess.run(["ocrmypdf", "-q", "--skip-text", "-l", "eng+por", str(ent), str(sai)], capture_output=True)
        return sai.read_bytes() if p.returncode == 0 and sai.exists() else None


def baixar_planilha_google(sess, url, destino, prefixo):
    """Baixa uma planilha Google (link público) como XLSX e grava uma CSV por aba."""
    import pandas as pd
    m = re.search(r"/spreadsheets/(?:u/\d+/)?d/(e/)?([\w-]+)", url)
    if not m: return []
    if m.group(1):
        u = f"https://docs.google.com/spreadsheets/d/e/{m.group(2)}/pub?output=xlsx"
    else:
        u = f"https://docs.google.com/spreadsheets/d/{m.group(2)}/export?format=xlsx"
    r = sess.get(u, timeout=90); time.sleep(PAUSA)
    if r.status_code != 200 or r.content[:2] != b"PK": return []
    arqs = []
    for nome, df in pd.read_excel(io.BytesIO(r.content), sheet_name=None).items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        arq = destino / f"{prefixo}__{re.sub(r'[^A-Za-z0-9]+', '_', str(nome))[:40]}.csv"
        df.to_csv(arq, index=False); arqs.append(arq)
    return arqs


def pdf_para_paginas(dados):
    p = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", "-", "-"], input=dados, capture_output=True)
    if p.returncode != 0: raise RuntimeError(p.stderr.decode(errors="ignore")[:300])
    paginas = normalizar(p.stdout.decode("utf-8", errors="replace")).split("\f")
    if paginas and not paginas[-1].strip(): paginas = paginas[:-1]
    return paginas


def docx_para_texto(dados):
    import docx
    d = docx.Document(io.BytesIO(dados))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows: partes.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(partes)


def sheet_para_csvs(html, destino, texto_id):
    import pandas as pd
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    nomes = [li.get_text(strip=True) for li in soup.select("#sheet-menu li")] or []
    tabelas = pd.read_html(io.StringIO(html), flavor="lxml")
    arquivos, resumo = [], []
    for i, t in enumerate(tabelas):
        nome = nomes[i] if i < len(nomes) else f"aba{i+1}"
        t = t.dropna(how="all").dropna(axis=1, how="all")
        arq = destino / f"{texto_id}__{re.sub(r'[^A-Za-z0-9]+', '_', nome)[:40]}.csv"
        t.to_csv(arq, index=False); arquivos.append(arq)
        resumo.append(f"## Aba: {nome}\n\nArquivo: `{arq.as_posix()}` — {len(t)} linhas × {t.shape[1]} colunas\n\n" + t.head(15).to_markdown(index=False))
    return arquivos, "\n\n".join(resumo)


def cabecalho(doc, meta):
    linhas = ["---"]
    for k in ["titulo", "id", "jogo", "competicao", "temporada", "versao", "tipo_documento", "status_biblioteca", "url"]:
        linhas.append(f'{k}: "{str(doc.get(k, "")).replace(chr(34), chr(39))}"')
    for k in ["sha256", "paginas", "idioma", "baixado_em"]:
        linhas.append(f'{k}: "{meta.get(k, "")}"')
    if doc.get("mesmo_arquivo_que"): linhas.append(f'tambem_listado_como: "{doc["mesmo_arquivo_que"]}"')
    linhas += ["fonte: \"Riot Games – Biblioteca de Operações Competitivas (competitiveops.riotgames.com)\"", "---", ""]
    return "\n".join(linhas)


# ---------------------------------------------------------------- principal
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raiz", default=".")
    ap.add_argument("--forcar", action="store_true")
    ap.add_argument("--id")
    ap.add_argument("--local", help="pasta com arquivos já baixados (nome do arquivo = id ou chave_arquivo)")
    ap.add_argument("--salvar-originais", action="store_true")
    a = ap.parse_args()
    raiz = Path(a.raiz)
    for d in ["textos", "dados", "originais"]: (raiz / d).mkdir(exist_ok=True)

    docs = list(csv.DictReader(open(raiz / "catalogo" / "catalogo.csv", encoding="utf-8")))
    manifest_path = raiz / "manifest.csv"
    anterior = {r["texto_id"]: r for r in csv.DictReader(open(manifest_path, encoding="utf-8"))} if manifest_path.exists() else {}

    unicos = {}
    for d in docs:
        tid = Path(d["texto"]).stem
        unicos.setdefault(tid, d)
    if a.id:
        unicos = {k: v for k, v in unicos.items() if k == a.id or v["id"] == a.id}

    sess = requests.Session(); sess.headers["User-Agent"] = UA
    manifest, secoes = dict(anterior), []
    secoes_antigas = []
    if (raiz / "secoes.csv").exists():
        secoes_antigas = list(csv.DictReader(open(raiz / "secoes.csv", encoding="utf-8")))

    feitos = set()
    for tid, d in unicos.items():
        prev = anterior.get(tid)
        dl, modo = url_download(d["url"])
        imutavel = modo in ("pdf", "docx")  # URLs do CDN Sanity contêm o hash do conteúdo
        if (prev and prev["status"] == "ok" and prev["url"] == d["url"] and imutavel and "pouco texto" not in prev.get("detalhe", "")
                and not a.forcar and not a.local):
            continue
        reg = {k: "" for k in CAMPOS_MANIFEST}
        reg.update(texto_id=tid, chave_arquivo=d["chave_arquivo"], titulo=d["titulo"], url=d["url"],
                   url_download=dl or "", baixado_em=agora(), arquivo_texto=f"textos/{tid}.md")
        try:
            if modo == "indisponivel":
                reg.update(status="indisponivel", detalhe="link interno da Riot, não público"); manifest[tid] = reg; continue
            dados = None
            if a.local:
                cands = [p for p in Path(a.local).glob("*") if p.stem in (tid, d["id"], d["chave_arquivo"]) or d["chave_arquivo"] in p.name]
                if cands: dados = cands[0].read_bytes(); reg["detalhe"] = f"arquivo local {cands[0].name}"
                elif a.local: continue
            if dados is None:
                r = sess.get(dl, timeout=90, allow_redirects=True); reg["http"] = r.status_code
                r.raise_for_status(); dados = r.content
                time.sleep(PAUSA)
                if modo == "drive" and b"%PDF" not in dados[:1024] and b"confirm=" in dados:
                    m = re.search(rb'href="(/uc\?export=download[^"]+)"', dados) or re.search(rb'action="([^"]+)"', dados)
                    if m:
                        u2 = m.group(1).decode().replace("&amp;", "&")
                        r = sess.get(u2 if u2.startswith("http") else "https://drive.google.com" + u2, timeout=90); dados = r.content
            reg["sha256"] = hashlib.sha256(dados).hexdigest(); reg["bytes"] = len(dados)

            corpo, paginas = "", []
            links = []
            if dados[:4] == b"%PDF":
                paginas = pdf_para_paginas(dados)
                n_pal = len(re.findall(r"\w+", "\n".join(paginas)))
                if paginas and n_pal < 10 * len(paginas):  # provável PDF-imagem
                    com_ocr = ocr(dados)
                    if com_ocr:
                        paginas = pdf_para_paginas(com_ocr); reg["detalhe"] = "texto obtido por OCR (conferir)"
                corpo = "\n".join(f"\n<!-- página {i} -->\n{pg.rstrip()}" for i, pg in enumerate(paginas, 1))
                links = links_pdf(dados)
                planilhas = [u for u in links if "docs.google.com/spreadsheets" in u]
                for j, u in enumerate(planilhas):
                    try:
                        arqs = baixar_planilha_google(sess, u, raiz / "dados", tid if j == 0 else f"{tid}_{j+1}")
                    except Exception as e2:
                        arqs = []; print(f"  planilha vinculada não baixada: {str(e2)[:120]}")
                    if arqs:
                        reg["detalhe"] = (reg["detalhe"] + "; " if reg["detalhe"] else "") + f"planilha vinculada exportada ({len(arqs)} abas em dados/)"
                if links:
                    corpo += "\n\n## Links presentes no documento\n\n" + "\n".join(f"- {u}" for u in links)
                ext = "pdf"
            elif modo == "docx" or dados[:2] == b"PK":
                corpo = normalizar(docx_para_texto(dados)); ext = "docx"
            elif modo == "gsheet":
                arqs, corpo = sheet_para_csvs(dados.decode("utf-8", errors="replace"), raiz / "dados", tid)
                reg["detalhe"] = f"{len(arqs)} abas exportadas para dados/"; ext = "html"
            elif modo == "gdoc_pub" and b"<html" in dados[:500].lower():
                from bs4 import BeautifulSoup
                corpo = normalizar(BeautifulSoup(dados, "html.parser").get_text("\n")); ext = "html"
            else:
                corpo = normalizar(dados.decode("utf-8", errors="replace")); ext = "txt"
                if "<html" in corpo[:500].lower():
                    raise RuntimeError("recebido HTML em vez do documento (link pode exigir login)")

            if a.salvar_originais:
                (raiz / "originais" / f"{tid}.{ext}").write_bytes(dados)
            texto_plano = "\n".join(paginas) if paginas else corpo
            reg["paginas"] = len(paginas) or ""
            reg["palavras"] = len(re.findall(r"\w+", texto_plano))
            reg["idioma"] = detectar_idioma(texto_plano)
            if paginas and reg["palavras"] < 10 * len(paginas):
                reg["detalhe"] = (reg["detalhe"] + "; " if reg["detalhe"] else "") + "pouco texto: PDF pode ser imagem (precisa OCR)"
            (raiz / "textos" / f"{tid}.md").write_text(cabecalho(d, reg) + corpo.strip() + "\n", encoding="utf-8")
            for num, tit, pg in indexar_secoes(paginas):
                secoes.append({"texto_id": tid, "secao": num, "titulo_secao": tit, "pagina": pg})
            reg["status"] = "ok"; feitos.add(tid)
        except Exception as e:  # registra a falha e segue
            msg = str(e)
            if "404" in msg and "sanity" in msg: msg = "link quebrado na própria Biblioteca (HTTP 404). " + msg
            elif "401" in msg or "403" in msg: msg = "documento Google com acesso restrito (não público). " + msg
            reg["status"] = "erro"; reg["detalhe"] = msg[:300]
        manifest[tid] = reg
        print(f"[{reg['status']}] {tid} {reg['detalhe']}", flush=True)

    with open(manifest_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CAMPOS_MANIFEST); w.writeheader()
        for k in sorted(manifest): w.writerow({c: manifest[k].get(c, "") for c in CAMPOS_MANIFEST})
    todas = [s for s in secoes_antigas if s["texto_id"] not in feitos] + secoes
    with open(raiz / "secoes.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["texto_id", "secao", "titulo_secao", "pagina"]); w.writeheader()
        w.writerows(sorted(todas, key=lambda s: (s["texto_id"], int(s["pagina"]))))
    st = [m["status"] for m in manifest.values()]
    print(json.dumps({s: st.count(s) for s in set(st)}, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
